"""Reconcile the NL invoice export with RepairDesk's current invoice statuses.

The workbook calls some historical discounts a Due Amount. Its refund-related
payment rows also differ from the current invoice list. Keep the raw export in
the source payload while importing the actual invoice totals and receivables.
"""

from __future__ import annotations

import argparse
import json
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


CENT = Decimal("0.01")


def amount(value: object) -> Decimal:
    return Decimal(str(value or "0")).quantize(CENT, rounding=ROUND_HALF_UP)


def currency(value: Decimal) -> str:
    return format(value.quantize(CENT, rounding=ROUND_HALF_UP), ".2f")


# Read from the live RepairDesk invoice list on 2026-10-02. These 20 invoices
# make up its entire Partial + Unpaid filter and its $1,636.95 receivable.
PENDING_DUE = {
    4912: "180.00", 4799: "289.00", 4756: "50.00", 4580: "49.00",
    4497: "75.00", 4385: "40.00", 4334: "59.00", 4137: "241.00",
    4086: "25.00", 3749: "29.95", 3318: "2.00", 3289: "5.00",
    3075: "24.00", 3024: "12.00", 2945: "54.00", 2879: "69.00",
    2526: "209.00", 2503: "60.00", 1810: "114.00", 169: "50.00",
}

# RepairDesk's Refunded filter contains 28 negative refund invoices and 40
# linked or adjusted invoices. The workbook has negative payment figures on
# some of the latter even though the current invoice row is settled.
REFUND_FILTER_IDS = {
    4873, 4836, 4818, 4804, 4682, 4542, 4517, 4427, 4289, 4285,
    4158, 4125, 4023, 3942, 3909, 3846, 3812, 3778, 3777, 3541,
    2991, 2906, 2839, 2563, 2375, 2342, 2283, 2275, 2241, 2226,
    2148, 2030, 1974, 1960, 1906, 1748, 1711, 1692, 1687, 1539,
    1485, 1476, 1370, 1317, 1131, 1049, 1033, 994, 968, 914,
    895, 890, 867, 850, 814, 695, 628, 606, 553, 430, 363, 346,
    345, 184, 173, 172, 171, 67,
}


def export_lines(path: Path) -> dict[int, list[dict]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").strip() for value in next(rows)]
    grouped: dict[int, list[dict]] = defaultdict(list)
    for values in rows:
        row = dict(zip(headers, values))
        if str(row.get("Store name") or "").strip().lower() != "techm8 north lakes":
            continue
        raw_number = str(row.get("Invoice number") or "").strip()
        if raw_number.isdecimal():
            grouped[int(raw_number)].append(row)
    workbook.close()
    return grouped


def write_batches(invoices: list[dict], output_dir: Path, batch_size: int) -> None:
    for stale in output_dir.glob("batch-*.sql"):
        stale.unlink()
    for index, offset in enumerate(range(0, len(invoices), batch_size), 1):
        value = json.dumps(invoices[offset:offset + batch_size], ensure_ascii=False, separators=(",", ":"))
        sql = "select public.import_repairdesk_sales_batch('" + value.replace("'", "''") + "'::jsonb);\n"
        (output_dir / f"batch-{index:03d}.sql").write_text(sql, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--prepared", required=True, type=Path)
    parser.add_argument("--invoice-export", required=True, type=Path)
    parser.add_argument("--batch-size", type=int, default=50)
    args = parser.parse_args()

    invoices = json.loads(args.prepared.read_text(encoding="utf-8"))
    if any("source_export_total" in row.get("order_payload", {}) for row in invoices):
        raise SystemExit("Prepared data is already reconciled; rebuild it from the workbook first")
    workbook_rows = export_lines(args.invoice_export)
    numbers = {row["invoice_number"] for row in invoices}
    assert len(invoices) == len(numbers) == len(workbook_rows) == 4914
    assert PENDING_DUE.keys() <= numbers and REFUND_FILTER_IDS <= numbers
    assert not (PENDING_DUE.keys() & REFUND_FILTER_IDS)

    counts = defaultdict(int)
    for invoice in invoices:
        number = invoice["invoice_number"]
        rows = workbook_rows[number]
        assert len(rows) == len(invoice["items"]), number
        gross = amount(invoice["total"])
        export_paid = amount(invoice["amount_paid"])
        source = invoice["order_payload"]
        source["source_export_total"] = currency(gross)
        source["source_export_amount_paid"] = currency(export_paid)
        source["source_export_payments"] = invoice["payments"]

        for line, row in zip(invoice["items"], rows):
            line_payload = line["line_payload"]
            line_payload["source_export_gross"] = currency(amount(row.get("Sub Total")))
            line_payload["source_export_amount_paid"] = currency(amount(row.get("Amount Paid")))
            line_payload["source_export_due_amount"] = currency(amount(row.get("Due Amount")))
            if line_payload.get("source_ticket_id") and not line["category"]:
                line["line_type"] = "repair"

        if number in PENDING_DUE:
            due = amount(PENDING_DUE[number])
            assert gross - export_paid == due, number
            invoice["payment_status"] = "deposit"
            source["source_invoice_status"] = "Unpaid" if export_paid == 0 else "Partial"
            source["source_current_due"] = currency(due)
            source["source_import_reconciliation"] = "outstanding"
            counts["outstanding"] += 1
        elif number in REFUND_FILTER_IDS:
            # All 68 current rows have no positive balance. One unusual refund
            # (#867) displays Paid $0 against a -$65 invoice.
            live_paid = Decimal("0") if number == 867 else gross
            invoice["amount_paid"] = currency(live_paid)
            invoice["payment_status"] = "paid"
            source["source_invoice_status"] = "Refunded" if gross < 0 else "Paid"
            source["source_current_due"] = "0.00"
            source["source_import_reconciliation"] = "refund_related"
            if live_paid != export_paid:
                invoice["payments"] = ([{
                    "method": "RepairDesk historical settlement",
                    "amount": currency(live_paid),
                    "taken_at": invoice["created_at"],
                }] if live_paid else [])
            counts["refund_related"] += 1
        elif export_paid < gross:
            # The current RepairDesk list reports these as Paid with zero Due.
            # Export Amount Paid is the net line amount after the discount.
            net_lines = [amount(row.get("Amount Paid")) for row in rows]
            assert sum(net_lines) == export_paid and all(x >= 0 for x in net_lines), number
            for line, row, net in zip(invoice["items"], rows, net_lines):
                original = amount(line["line_total"])
                quantity = int(line["quantity"])
                line["line_total"] = currency(net)
                line["unit_price"] = format((net / quantity).quantize(Decimal("0.000001"), rounding=ROUND_HALF_UP), ".6f")
                meta = line["line_payload"]
                meta["source_discount"] = currency(original - net)
                meta["source_export_tax"] = meta["source_tax"]
                tax = (net / Decimal("11")).quantize(CENT, rounding=ROUND_HALF_UP) if amount(row.get("Tax")) else Decimal("0")
                meta["source_tax"] = currency(tax)
                meta["source_total_sales_ex_gst"] = currency(net - tax)
                assert amount(line["unit_price"]) * quantity == net or amount(Decimal(line["unit_price"]) * quantity) == net, number
            invoice["total"] = currency(export_paid)
            invoice["payment_status"] = "paid"
            source["source_invoice_status"] = "Paid"
            source["source_current_due"] = "0.00"
            source["source_import_reconciliation"] = "discount_from_export_due"
            counts["discounts"] += 1
        else:
            assert export_paid == gross, number
            invoice["payment_status"] = "paid"
            source["source_invoice_status"] = "Paid"
            source["source_current_due"] = "0.00"
            source["source_import_reconciliation"] = "as_exported"
            counts["as_exported"] += 1

        assert sum(amount(item["line_total"]) for item in invoice["items"]) == amount(invoice["total"]), number
        assert sum(amount(payment["amount"]) for payment in invoice["payments"]) == amount(invoice["amount_paid"]), number

    assert dict(counts) == {"outstanding": 20, "refund_related": 68, "discounts": 67, "as_exported": 4759}, counts
    assert sum(amount(PENDING_DUE[number]) for number in PENDING_DUE) == Decimal("1636.95")
    args.prepared.write_text(json.dumps(invoices, ensure_ascii=False, separators=(",", ":")) + "\n", encoding="utf-8")
    write_batches(invoices, args.prepared.parent, args.batch_size)
    result = {
        "invoice_count": len(invoices),
        "line_count": sum(len(invoice["items"]) for invoice in invoices),
        "payment_count": sum(len(invoice["payments"]) for invoice in invoices),
        "total": currency(sum(amount(invoice["total"]) for invoice in invoices)),
        "amount_paid": currency(sum(amount(invoice["amount_paid"]) for invoice in invoices)),
        "outstanding_balance": currency(sum(amount(invoice["total"]) - amount(invoice["amount_paid"]) for invoice in invoices if invoice["payment_status"] == "deposit")),
        "classification_counts": dict(counts),
    }
    (args.prepared.parent / "reconciliation.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
